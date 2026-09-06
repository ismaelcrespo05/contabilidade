import re
from datetime import datetime

NOMES_MESES = {
    "janeiro": 1, "fevereiro": 2, "março": 3, "marco": 3, "abril": 4,
    "maio": 5, "junho": 6, "julho": 7, "agosto": 8, "setembro": 9,
    "outubro": 10, "novembro": 11, "dezembro": 12,
}

LABELS_IGNORADOS = {"rotinas fiscais"}


def _extrair_regime_do_titulo(aba):
    for row in aba.iter_rows(min_row=1, max_row=6):
        for cell in row:
            if isinstance(cell.value, str) and "ACOMPANHAMENTO" in cell.value.upper():
                m = re.search(r"Empresas\s+(.+?)\.?\s*$", cell.value.strip())
                if m:
                    return m.group(1).strip().rstrip(".")
    return None


def _cor_preenchimento(aba, linha, coluna):
    celula = aba.cell(row=linha, column=coluna)
    try:
        return (celula.fill.fgColor.rgb or "").upper()
    except Exception:
        return ""


def _linhas_mescladas_largura_total(aba):
    """Retorna {linha: (col_nome, ultima_col_mes)} só para linhas mescladas
    que sejam realmente nome de empresa ou divisor "Rotinas Fiscais" —
    identificadas pelo preenchimento cinza padrão (FFE7E7E7). Sub-títulos
    mesclados com outra cor (ex: amarelo) são ignorados aqui de propósito."""
    resultado = {}
    for rng in aba.merged_cells.ranges:
        if rng.min_row == rng.max_row and (rng.max_col - rng.min_col) >= 10:
            cor = _cor_preenchimento(aba, rng.min_row, rng.min_col)
            if cor in ("FFE7E7E7", "E7E7E7"):
                resultado[rng.min_row] = (rng.min_col, rng.max_col)
    return resultado


def processar_planilha(caminho, ano_padrao):
    """
    Lê o arquivo (caminho no disco ou arquivo já aberto, tipo
    request.FILES["arquivo"]) e devolve uma lista de eventos:
      ("empresa", nome_do_regime, nome_da_empresa)
      ("obrigacao", nome_da_empresa, nome_da_obrigacao, [(mes, ano, status), ...])

    Entende os dois estilos de planilha que aparecem no arquivo modelo:
    cabeçalho de meses em texto sem ano (usa ano_padrao) e cabeçalho com
    datas reais por empresa (usa o ano de cada data, quando disponível).
    """
    import openpyxl
    workbook = openpyxl.load_workbook(caminho, data_only=True)

    log = []

    for aba in workbook.worksheets:
        if not aba.title.strip().upper().startswith("DEPT. FISCAL"):
            continue

        regime_nome = _extrair_regime_do_titulo(aba)
        if not regime_nome:
            continue

        linhas_full = _linhas_mescladas_largura_total(aba)

        empresa_atual = None
        col_nome = None
        mapa_colunas_mes = {}

        max_row = aba.max_row
        for num_linha in range(1, max_row + 1):

            if num_linha in linhas_full:
                c_nome, c_ultima = linhas_full[num_linha]
                valor = aba.cell(row=num_linha, column=c_nome).value
                texto = str(valor).strip() if valor else ""

                if "ACOMPANHAMENTO" in texto.upper():
                    continue
                if texto.lower() in LABELS_IGNORADOS:
                    continue

                empresa_atual = texto
                col_nome = c_nome
                log.append(("empresa", regime_nome, empresa_atual))
                continue

            if col_nome is None:
                continue

            valor_nome_col = aba.cell(row=num_linha, column=col_nome).value

            if isinstance(valor_nome_col, str) and valor_nome_col.strip().lower() == "rotinas fiscais":
                mapa_colunas_mes = {}
                for col in range(col_nome + 1, col_nome + 13):
                    v = aba.cell(row=num_linha, column=col).value
                    if isinstance(v, datetime):
                        mapa_colunas_mes[col] = (v.month, v.year)
                continue

            if not empresa_atual or not valor_nome_col:
                continue

            nome_obrigacao = str(valor_nome_col).strip()
            if nome_obrigacao.lower() in NOMES_MESES:
                continue

            valores = []
            if mapa_colunas_mes:
                for col, (mes, ano) in mapa_colunas_mes.items():
                    v = aba.cell(row=num_linha, column=col).value
                    if isinstance(v, (int, float)) and int(v) in (1, 2, 3):
                        valores.append((mes, ano, int(v)))
            else:
                for col in range(col_nome + 1, col_nome + 13):
                    v = aba.cell(row=num_linha, column=col).value
                    if isinstance(v, (int, float)) and int(v) in (1, 2, 3):
                        valores.append((col - col_nome, ano_padrao, int(v)))

            if valores:
                log.append(("obrigacao", empresa_atual, nome_obrigacao, valores))

    return log


def importar_log_para_banco(log):
    """
    Recebe o "log" de processar_planilha() e grava tudo no banco: cria/
    reaproveita RegimeTributario e Empresa (só com o nome, se ainda não
    existir), ObrigacaoAcessoria, Competencia e CumprimentoObrigacao
    (status gravado em status_manual). Idempotente: rodar de novo com a
    mesma planilha não duplica nada.
    """
    from .models import (
        RegimeTributario, Empresa, ObrigacaoAcessoria, Competencia,
        CumprimentoObrigacao,
    )

    empresas_criadas = 0
    empresas_existentes = 0
    obrigacoes_criadas = 0
    cumprimentos_gravados = 0

    cache_empresas = {}

    for item in log:
        if item[0] == "empresa":
            _, regime_nome, nome_empresa = item

            regime, _ = RegimeTributario.objects.get_or_create(nome=regime_nome)

            empresa, criada = Empresa.objects.get_or_create(
                razao_social=nome_empresa,
                regime_tributario=regime,
            )
            cache_empresas[nome_empresa] = empresa

            if criada:
                empresas_criadas += 1
            else:
                empresas_existentes += 1

    for item in log:
        if item[0] == "obrigacao":
            _, nome_empresa, nome_obrigacao, valores = item

            empresa = cache_empresas.get(nome_empresa)
            if not empresa:
                continue

            obrigacao, criada = ObrigacaoAcessoria.objects.get_or_create(nome=nome_obrigacao)
            if criada:
                obrigacoes_criadas += 1

            if obrigacao not in empresa.obrigacoes_acessorias.all():
                empresa.obrigacoes_acessorias.add(obrigacao)

            for mes, ano, status in valores:
                competencia, _ = Competencia.objects.get_or_create(
                    empresa=empresa, ano=ano, mes=mes
                )
                cumprimento, _ = CumprimentoObrigacao.objects.get_or_create(
                    competencia=competencia, obrigacao=obrigacao
                )
                cumprimento.status_manual = status
                cumprimento.save()
                cumprimentos_gravados += 1

    return {
        "empresas_criadas": empresas_criadas,
        "empresas_existentes": empresas_existentes,
        "obrigacoes_criadas": obrigacoes_criadas,
        "cumprimentos_gravados": cumprimentos_gravados,
    }