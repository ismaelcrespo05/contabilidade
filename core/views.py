from django.contrib import messages
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse
from django.views.generic import ListView, DeleteView
from django.contrib.auth import views as auth_views
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from .forms import EmpresaForm, CompetenciaForm, ApuracaoImpostoForm, DocumentoForm, EditarUsuarioForm, ImportarRotinasForm, NovoUsuarioForm, LancamentoContabilForm, CertificadoForm, CertificadoRapidoFormSet,  ConfiguracaoSistemaForm
from .models import (
    TipoEmpresa, RegimeTributario, Imposto, CNAE, ObrigacaoAcessoria,
    Empresa, Competencia, ApuracaoImposto, CumprimentoObrigacao, Documento, Fornecedor, LancamentoContabil, ConfiguracaoSistema, Certificado, Notificacao
    )

from django.contrib.auth.models import User

from .permissoes import (
    requer_edicao, requer_exclusao, requer_administrador,
    RequerEdicaoMixin, RequerExclusaoMixin, pode_editar,
)



def inicio(request):
    return render(request, "core/inicio.html")


# ---------------------------------------------------------------------
# Login / Logout personalizados (substituem o login do admin do Django)
# ---------------------------------------------------------------------
class PainelLoginView(auth_views.LoginView):
    template_name = "core/login.html"
    redirect_authenticated_user = True


class PainelLogoutView(auth_views.LogoutView):
    next_page = "login"


# ---------------------------------------------------------------------
# Registro de catálogos administrados pelo painel.
# ---------------------------------------------------------------------
CATALOGOS = {
    "tipos-empresa": {
        "model": TipoEmpresa, "titulo": "Tipos de Empresa",
        "titulo_singular": "Tipo de Empresa", "fields": ["nome"],
        "columns": [("nome", "Nome")],
    },
    "regimes": {
        "model": RegimeTributario, "titulo": "Regimes Tributários",
        "titulo_singular": "Regime Tributário", "fields": ["nome"],
        "columns": [("nome", "Nome")],
    },
    "impostos": {
        "model": Imposto, "titulo": "Impostos",
        "titulo_singular": "Imposto", "fields": ["nome", "descricao"],
        "columns": [("nome", "Nome"), ("descricao", "Descrição")],
    },
    "cnae": {
        "model": CNAE, "titulo": "CNAE",
        "titulo_singular": "CNAE", "fields": ["codigo", "descricao"],
        "columns": [("codigo", "Código"), ("descricao", "Descrição")],
    },
    "obrigacoes": {
        "model": ObrigacaoAcessoria, "titulo": "Obrigações Acessórias",
        "titulo_singular": "Obrigação Acessória", "fields": ["nome", "descricao", "dia_vencimento"],
        "columns": [("nome", "Nome"), ("descricao", "Descrição"), ("dia_vencimento", "Vence dia")],
    },
        "fornecedores": {
        "model": Fornecedor, "titulo": "Fornecedores",
        "titulo_singular": "Fornecedor", "fields": ["cnpj", "razao_social", "nome_fantasia"],
        "columns": [("cnpj", "CNPJ"), ("razao_social", "Razão Social"), ("nome_fantasia", "Nome Fantasia")],
    },
}


def get_catalogo_ou_404(slug):
    if slug not in CATALOGOS:
        from django.http import Http404
        raise Http404("Catálogo não encontrado")
    return CATALOGOS[slug]


@login_required
def painel_dashboard(request):
    context = {
        "total_empresas": Empresa.objects.count(),
        "total_impostos": Imposto.objects.count(),
        "apuracoes_pendentes": ApuracaoImposto.objects.filter(
            status__in=["PENDENTE", "ATRASADO"]
        ).count(),
        "catalogos": [
            {"slug": slug, "titulo": conf["titulo"]}
            for slug, conf in CATALOGOS.items()
        ],
    }
    return render(request, "core/painel_dashboard.html", context)


# ---------------------------------------------------------------------
# Empresas
# ---------------------------------------------------------------------
class EmpresaListView(LoginRequiredMixin, ListView):
    model = Empresa
    template_name = "core/empresa_list.html"
    context_object_name = "empresas"
    paginate_by = 20

    def get_queryset(self):
        from .list_utils import aplicar_busca_e_ordenacao

        qs, self.termo, self.ordenar = aplicar_busca_e_ordenacao(
            self.request,
            Empresa.objects.select_related("tipo_empresa", "regime_tributario"),
            campos_busca=["razao_social", "cnpj"],
            opcoes_ordenacao={"nome": "razao_social", "cnpj": "cnpj"},
            padrao="nome",
        )
        return qs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["termo"] = self.termo
        context["ordenar"] = self.ordenar
        return context


@requer_edicao
def nova_empresa(request):
    # Se vier de um Documento sem empresa cadastrada (RF-13), o CNPJ/razão
    # social extraídos chegam por querystring, e ao salvar a empresa nova
    # ela já fica vinculada de volta a esse documento.
    documento_id = request.GET.get("documento_id") or request.POST.get("documento_id")

    if request.method == "POST":
        form = EmpresaForm(request.POST)
        if form.is_valid():
            empresa = form.save()
            messages.success(request, "Empresa cadastrada com sucesso.")

            if documento_id:
                documento = Documento.objects.filter(pk=documento_id).first()
                if documento:
                    documento.empresa = empresa
                    documento.save()
                    return redirect("revisar_documento", pk=documento.pk)

            return redirect("empresa_list")
    else:
        initial = {}
        if request.GET.get("cnpj"):
            initial["cnpj"] = request.GET["cnpj"]
        if request.GET.get("razao_social"):
            initial["razao_social"] = request.GET["razao_social"]
        form = EmpresaForm(initial=initial)

    return render(request, "core/empresa_form.html", {
        "form": form, "titulo": "Nova Empresa", "documento_id": documento_id
    })

@requer_edicao
def editar_empresa(request, pk):
    empresa = get_object_or_404(Empresa, pk=pk)
    if request.method == "POST":
        form = EmpresaForm(request.POST, instance=empresa)
        if form.is_valid():
            form.save()
            messages.success(request, "Empresa atualizada com sucesso.")
            return redirect("empresa_list")
    else:
        form = EmpresaForm(instance=empresa)

    return render(request, "core/empresa_form.html", {
        "form": form, "titulo": f"Editar — {empresa.razao_social}"
    })


class EmpresaDeleteView(RequerExclusaoMixin, DeleteView):
    model = Empresa
    template_name = "core/confirm_delete.html"
    success_url = "/painel/empresas/"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["cancel_url"] = self.success_url
        context["titulo"] = "Excluir empresa"
        return context


# ---------------------------------------------------------------------
# Catálogos
# ---------------------------------------------------------------------
@login_required
def catalogo_list(request, slug):
    conf = get_catalogo_ou_404(slug)
    itens = conf["model"].objects.all().order_by("pk")
    return render(request, "core/catalogo_list.html", {"slug": slug, "conf": conf, "itens": itens})


@requer_edicao
def catalogo_form(request, slug, pk=None):
    from django.forms import modelform_factory

    conf = get_catalogo_ou_404(slug)
    Form = modelform_factory(conf["model"], fields=conf["fields"])

    documento_id = request.GET.get("documento_id") or request.POST.get("documento_id")

    instancia = None
    if pk is not None:
        instancia = get_object_or_404(conf["model"], pk=pk)

    if request.method == "POST":
        form = Form(request.POST, instance=instancia)
        if form.is_valid():
            objeto = form.save()
            messages.success(request, "Registro salvo com sucesso.")

            if documento_id and slug == "fornecedores":
                documento = Documento.objects.filter(pk=documento_id).first()
                if documento:
                    documento.fornecedor = objeto
                    documento.save()
                    return redirect("revisar_documento", pk=documento.pk)

            return redirect("catalogo_list", slug=slug)
    else:
        initial = {}
        for campo in conf["fields"]:
            valor = request.GET.get(campo)
            if valor:
                initial[campo] = valor
        form = Form(instance=instancia, initial=initial)

    return render(request, "core/catalogo_form.html", {
        "form": form, "conf": conf, "slug": slug, "documento_id": documento_id,
        "titulo": (f"Editar {conf['titulo_singular']}" if instancia else f"Novo {conf['titulo_singular']}"),
    })


@requer_exclusao
def catalogo_delete(request, slug, pk):
    conf = get_catalogo_ou_404(slug)
    instancia = get_object_or_404(conf["model"], pk=pk)

    if request.method == "POST":
        instancia.delete()
        messages.success(request, "Registro excluído com sucesso.")
        return redirect("catalogo_list", slug=slug)

    return render(request, "core/confirm_delete.html", {
        "object": instancia,
        "cancel_url": reverse("catalogo_list", args=[slug]),
        "titulo": f"Excluir {conf['titulo_singular']}",
    })


# ---------------------------------------------------------------------
# Competências
# ---------------------------------------------------------------------
@login_required
def competencia_list(request, empresa_pk):
    empresa = get_object_or_404(Empresa, pk=empresa_pk)
    competencias = empresa.competencias.order_by("-ano", "-mes")
    return render(request, "core/competencia_list.html", {"empresa": empresa, "competencias": competencias})


@requer_edicao
def nova_competencia(request, empresa_pk):
    empresa = get_object_or_404(Empresa, pk=empresa_pk)

    if request.method == "POST":
        form = CompetenciaForm(request.POST)
        if form.is_valid():
            competencia = form.save(commit=False)
            competencia.empresa = empresa
            competencia.save()

            CumprimentoObrigacao.objects.bulk_create([
                CumprimentoObrigacao(competencia=competencia, obrigacao=obrigacao)
                for obrigacao in empresa.obrigacoes_acessorias.all()
            ])

            messages.success(request, "Competência criada com sucesso.")
            return redirect("competencia_list", empresa_pk=empresa.pk)
    else:
        form = CompetenciaForm()

    return render(request, "core/competencia_form.html", {"form": form, "empresa": empresa})


class CompetenciaDeleteView(RequerExclusaoMixin, DeleteView):
    model = Competencia
    template_name = "core/confirm_delete.html"

    def get_success_url(self):
        return reverse("competencia_list", args=[self.object.empresa_id])

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["cancel_url"] = self.get_success_url()
        context["titulo"] = "Excluir competência"
        return context


@login_required
def competencia_detail(request, pk):
    competencia = get_object_or_404(Competencia, pk=pk)
    apuracoes = competencia.apuracoes_imposto.select_related("imposto")
    cumprimentos = competencia.cumprimentos_obrigacao.select_related("obrigacao").order_by("obrigacao__nome")

    if request.method == "POST":
        if not pode_editar(request.user):
            messages.error(request, "Seu perfil (Consulta) não tem permissão para editar.")
            return redirect("competencia_detail", pk=pk)
        from django.utils import timezone

        for cumprimento in cumprimentos:
            # Cada linha manda dois campos próprios, identificados pelo id:
            # "concluido_<id>" (checkbox) e "status_manual_<id>" (select).
            concluido_agora = request.POST.get(f"concluido_{cumprimento.id}") == "on"
            status_manual_raw = request.POST.get(f"status_manual_{cumprimento.id}", "")

            mudou = False

            if concluido_agora and not cumprimento.cumprido:
                cumprimento.cumprido = True
                cumprimento.data_cumprimento = timezone.localdate()
                mudou = True
            elif not concluido_agora and cumprimento.cumprido:
                cumprimento.cumprido = False
                cumprimento.data_cumprimento = None
                mudou = True

            novo_status_manual = int(status_manual_raw) if status_manual_raw else None
            if novo_status_manual != cumprimento.status_manual:
                cumprimento.status_manual = novo_status_manual
                mudou = True

            if mudou:
                cumprimento.save()

        messages.success(request, "Rotinas fiscais atualizadas.")
        return redirect("competencia_detail", pk=competencia.pk)

    return render(request, "core/competencia_detail.html", {
        "competencia": competencia, "apuracoes": apuracoes, "cumprimentos": cumprimentos,
    })


# ---------------------------------------------------------------------
# Apurações de imposto
# ---------------------------------------------------------------------
@requer_edicao
def nova_apuracao(request, competencia_pk):
    competencia = get_object_or_404(Competencia, pk=competencia_pk)

    if request.method == "POST":
        form = ApuracaoImpostoForm(request.POST)
        if form.is_valid():
            apuracao = form.save(commit=False)
            apuracao.competencia = competencia
            apuracao.save()
            messages.success(request, "Apuração registrada com sucesso.")
            return redirect("competencia_detail", pk=competencia.pk)
    else:
        form = ApuracaoImpostoForm()

    return render(request, "core/apuracao_form.html", {
        "form": form, "competencia": competencia, "titulo": "Nova Apuração"
    })


@requer_edicao
def editar_apuracao(request, pk):
    apuracao = get_object_or_404(ApuracaoImposto, pk=pk)

    if request.method == "POST":
        form = ApuracaoImpostoForm(request.POST, instance=apuracao)
        if form.is_valid():
            form.save()
            messages.success(request, "Apuração atualizada com sucesso.")
            return redirect("competencia_detail", pk=apuracao.competencia_id)
    else:
        form = ApuracaoImpostoForm(instance=apuracao)

    return render(request, "core/apuracao_form.html", {
        "form": form, "competencia": apuracao.competencia, "titulo": "Editar Apuração",
    })

@login_required
def exportar_excel(request):
    """
    Gera um .xlsx no mesmo layout da planilha modelo enviada pelo cliente:
    título mesclado no topo, uma aba por regime tributário, um bloco por
    empresa (nome mesclado + linha "Rotinas Fiscais" mesclada + uma linha
    por obrigação), sem linha em branco entre empresas — e o "semáforo"
    (verde/amarelo/vermelho) em cada mês usando o recurso real de
    formatação condicional do Excel (iconSet), não um texto colorido.

    Valores gravados na célula (ocultos, só o ícone aparece):
      3 = Em dia (verde) | 2 = Risco de atraso (amarelo) | 1 = Atrasada (vermelho)

    Parâmetros via querystring: ?ano=2026 (padrão: ano atual)
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.formatting.rule import IconSetRule
    from django.http import HttpResponse
    from datetime import date

    ano = int(request.GET.get("ano", date.today().year))
    meses_nomes = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
                   "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    titulo_font = Font(bold=True, size=18)
    secao_fill = PatternFill("solid", fgColor="E7E7E7")
    secao_font = Font(bold=True, size=12)
    centro = Alignment(horizontal="center")
    esquerda = Alignment(horizontal="left")

    def nova_regra_semaforo():
        return IconSetRule(
            icon_style="3TrafficLights1",
            type="num",
            values=[0, 2, 3],
            showValue=False,
        )

    for regime in RegimeTributario.objects.all():
        aba = workbook.create_sheet(title=regime.nome[:31])

        aba.merge_cells("B1:N1")
        titulo = aba.cell(row=1, column=2,
            value=f"ACOMPANHAMENTO: Departamento Fiscal - Empresas {regime.nome}.")
        titulo.font = titulo_font
        titulo.alignment = centro
        aba.row_dimensions[1].height = 31.5

        aba.cell(row=1, column=16, value="Legenda").fill = secao_fill
        aba.cell(row=2, column=16, value="Em dia (3)")
        aba.cell(row=2, column=17, value=3)
        aba.cell(row=3, column=16, value="Risco de atraso (2)")
        aba.cell(row=3, column=17, value=2)
        aba.cell(row=4, column=16, value="Atrasada (1)")
        aba.cell(row=4, column=17, value=1)
        aba.conditional_formatting.add("Q2:Q4", nova_regra_semaforo())

        for col, nome_mes in enumerate(meses_nomes, start=3):
            celula = aba.cell(row=2, column=col, value=nome_mes)
            celula.alignment = centro
            celula.font = Font(bold=True)
            celula.fill = secao_fill

        linha = 3
        empresas = Empresa.objects.filter(
            regime_tributario=regime
        ).prefetch_related("obrigacoes_acessorias")

        for empresa in empresas:
            aba.merge_cells(f"B{linha}:N{linha}")
            celula = aba.cell(row=linha, column=2, value=empresa.razao_social)
            celula.fill = secao_fill
            celula.font = secao_font
            celula.alignment = centro
            linha += 1

            obrigacoes = empresa.obrigacoes_acessorias.all()
            if not obrigacoes:
                continue

            aba.merge_cells(f"B{linha}:N{linha}")
            celula = aba.cell(row=linha, column=2, value="Rotinas Fiscais")
            celula.fill = secao_fill
            celula.font = secao_font
            celula.alignment = esquerda
            linha += 1

            cumprimentos = CumprimentoObrigacao.objects.filter(
                competencia__empresa=empresa, competencia__ano=ano,
            ).select_related("competencia", "obrigacao")

            por_mes_obrigacao = {
                (c.competencia.mes, c.obrigacao_id): c.status_calculado()
                for c in cumprimentos
            }

            for obrigacao in obrigacoes:
                linha_inicial = linha
                aba.cell(row=linha, column=2, value=obrigacao.nome)

                for mes in range(1, 13):
                    status = por_mes_obrigacao.get((mes, obrigacao.id))
                    if status is not None:
                        aba.cell(row=linha, column=2 + mes, value=status)

                aba.conditional_formatting.add(
                    f"C{linha_inicial}:N{linha_inicial}", nova_regra_semaforo()
                )
                linha += 1

            # Sem linha em branco entre empresas — igual à planilha modelo.

        aba.column_dimensions["A"].width = 3
        aba.column_dimensions["B"].width = 42
        for col in range(3, 15):
            aba.column_dimensions[
                openpyxl.utils.get_column_letter(col)
            ].width = 11
        aba.column_dimensions["P"].width = 21.5
        aba.column_dimensions["Q"].width = 8.4

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument"
            ".spreadsheetml.sheet"
        )
    )
    response["Content-Disposition"] = (
        f'attachment; filename="acompanhamento_fiscal_{ano}.xlsx"'
    )
    workbook.save(response)
    return response


class ApuracaoDeleteView(RequerExclusaoMixin, DeleteView):
    model = ApuracaoImposto
    template_name = "core/confirm_delete.html"

    def get_success_url(self):
        return reverse("competencia_detail", args=[self.object.competencia_id])

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["cancel_url"] = self.get_success_url()
        context["titulo"] = "Excluir apuração"
        return context
        
# ---------------------------------------------------------------------
# Importação de comprovante em PDF (recibo de entrega de SPED/EFD/etc.)
# ---------------------------------------------------------------------
@requer_edicao
def importar_comprovante(request):
    """
    Passo 1: o usuário sobe o PDF. O sistema lê o CNPJ e o período do
    recibo, encontra a empresa e a competência correspondentes (criando
    a competência se ainda não existir) e manda para a tela de
    confirmação — onde o usuário escolhe manualmente a qual obrigação
    esse recibo se refere.
    """
    if request.method == "POST" and request.FILES.get("arquivo"):
        from datetime import datetime
        from django.core.files.storage import default_storage
        from .pdf_utils import extrair_dados_recibo

        arquivo = request.FILES["arquivo"]

        try:
            dados = extrair_dados_recibo(arquivo)
        except Exception:
            messages.error(request, "Não consegui ler esse arquivo. Confirme que é um PDF válido.")
            return redirect("importar_comprovante")

        if not dados["cnpj"] or not dados["periodo_fim"]:
            messages.error(
                request,
                "Não consegui encontrar o CNPJ e/ou o período nesse PDF. "
                "Verifique se é um recibo de entrega válido."
            )
            return redirect("importar_comprovante")

        empresa = Empresa.objects.filter(cnpj=dados["cnpj"]).first()
        if not empresa:
            messages.error(
                request,
                f"Nenhuma empresa cadastrada com o CNPJ {dados['cnpj']}."
            )
            return redirect("importar_comprovante")

        data_fim = datetime.strptime(dados["periodo_fim"], "%d/%m/%Y").date()

        competencia, _ = Competencia.objects.get_or_create(
            empresa=empresa, ano=data_fim.year, mes=data_fim.month
        )

        existentes = set(
            competencia.cumprimentos_obrigacao.values_list("obrigacao_id", flat=True)
        )
        CumprimentoObrigacao.objects.bulk_create([
            CumprimentoObrigacao(competencia=competencia, obrigacao=obrigacao)
            for obrigacao in empresa.obrigacoes_acessorias.all()
            if obrigacao.id not in existentes
        ])

        arquivo.seek(0)
        caminho_temp = default_storage.save(
            f"comprovantes/_pendentes/{arquivo.name}", arquivo
        )

        request.session["comprovante_pendente"] = {
            "caminho": caminho_temp,
            "nome_arquivo": arquivo.name,
            "competencia_id": competencia.id,
            "cnpj": dados["cnpj"],
            "empresa_nome": empresa.razao_social,
            "periodo": f"{dados['periodo_inicio']} a {dados['periodo_fim']}",
        }

        return redirect("confirmar_comprovante")

    return render(request, "core/importar_comprovante.html")


@requer_edicao
def confirmar_comprovante(request):
    """
    Passo 2: mostra o que foi lido do PDF (empresa, período) e pede para
    o usuário escolher, num select, a qual obrigação esse recibo se
    refere. Ao confirmar, o PDF é vinculado a essa obrigação e ela é
    marcada como cumprida.
    """
    pendente = request.session.get("comprovante_pendente")

    if not pendente:
        messages.error(request, "Nenhum comprovante pendente de confirmação.")
        return redirect("importar_comprovante")

    competencia = get_object_or_404(Competencia, pk=pendente["competencia_id"])
    cumprimentos = competencia.cumprimentos_obrigacao.select_related("obrigacao").order_by("obrigacao__nome")

    if request.method == "POST":
        from django.core.files.storage import default_storage
        from django.core.files.base import ContentFile
        from django.utils import timezone

        cumprimento_id = request.POST.get("cumprimento_id")

        if not cumprimento_id:
            messages.error(request, "Escolha a qual obrigação esse comprovante se refere.")
            return redirect("confirmar_comprovante")

        cumprimento = get_object_or_404(CumprimentoObrigacao, pk=cumprimento_id)

        with default_storage.open(pendente["caminho"], "rb") as f:
            conteudo = f.read()

        cumprimento.comprovante.save(
            pendente["nome_arquivo"], ContentFile(conteudo), save=False
        )
        cumprimento.cumprido = True
        cumprimento.data_cumprimento = timezone.localdate()
        cumprimento.save()

        default_storage.delete(pendente["caminho"])
        del request.session["comprovante_pendente"]

        messages.success(
            request,
            f"Comprovante vinculado a \"{cumprimento.obrigacao.nome}\" e obrigação baixada com sucesso."
        )
        return redirect("competencia_detail", pk=competencia.pk)

    return render(request, "core/confirmar_comprovante.html", {
        "competencia": competencia,
        "cumprimentos": cumprimentos,
        "cnpj": pendente["cnpj"],
        "empresa_nome": pendente["empresa_nome"],
        "periodo": pendente["periodo"],
    })
    
# ---------------------------------------------------------------------
# Documentos fiscais genéricos (Fase 1 — RF-05 a RF-21)
# ---------------------------------------------------------------------
@login_required
def documento_list(request):
    documentos = Documento.objects.select_related("empresa").order_by("-data_upload")
    return render(request, "core/documento_list.html", {"documentos": documentos})


@requer_edicao
def novo_documento(request):
    """
    RF-06/RF-07: sobe o PDF. RF-08: extrai os dados automaticamente.
    RF-05: não exige empresa cadastrada para isso.
    """
    if request.method == "POST" and request.FILES.get("arquivo"):
        from .pdf_utils import extrair_dados_documento, parse_valor_brl, parse_data_br

        arquivo = request.FILES["arquivo"]

        try:
            dados = extrair_dados_documento(arquivo)
        except Exception:
            messages.error(
                request,
                "Não consegui processar esse arquivo. Confirme que é um PDF válido."
            )
            return redirect("novo_documento")

        documento = Documento.objects.create(
            arquivo=arquivo,
            cnpj_emitente=dados["cnpj_emitente"] or "",
            razao_social_emitente=dados["razao_social_emitente"],
            cnpj_destinatario=dados["cnpj_destinatario"] or "",
            numero_documento=dados["numero_documento"] or "",
            data_emissao=parse_data_br(dados["data_emissao"]),
            valor_total=parse_valor_brl(dados["valor_total"]),
            dados_extraidos={
                "cnpjs_candidatos": dados["cnpjs"],
                "razao_social_destinatario": dados["razao_social_destinatario"],
                "impostos": dados["impostos"],
                "texto": dados["texto"][:5000],
            },
        )

        # RF-10/RF-11/RF-12: procura empresa cadastrada pelo CNPJ do
        # emitente ou do destinatário. Se os dois existirem, deixa para o
        # usuário escolher na tela de revisão (RF-09: não presume sozinho
        # qual delas é a empresa administrada).
        empresa_emitente = (
            Empresa.objects.filter(cnpj=documento.cnpj_emitente).first()
            if documento.cnpj_emitente else None
        )
        empresa_destinatario = (
            Empresa.objects.filter(cnpj=documento.cnpj_destinatario).first()
            if documento.cnpj_destinatario else None
        )

        if empresa_emitente and not empresa_destinatario:
            documento.empresa = empresa_emitente
            documento.save()
        elif empresa_destinatario and not empresa_emitente:
            documento.empresa = empresa_destinatario
            documento.save()
        elif not empresa_emitente and not empresa_destinatario:
            messages.warning(
                request,
                "Empresa não cadastrada para o(s) CNPJ(s) encontrado(s) nesse "
                "documento. Você pode analisar sem cadastrar ou cadastrar agora."
            )
        # RF-31: tenta identificar um Fornecedor cadastrado pelo CNPJ do
        # emitente — em regra, quem emite uma nota para a empresa
        # administrada é o fornecedor.
        if documento.cnpj_emitente and documento.cnpj_emitente != getattr(documento.empresa, "cnpj", None):
            fornecedor = Fornecedor.objects.filter(cnpj=documento.cnpj_emitente).first()
            if fornecedor:
                documento.fornecedor = fornecedor
                documento.save()    

        return redirect("revisar_documento", pk=documento.pk)

    return render(request, "core/novo_documento.html")


@login_required
def revisar_documento(request, pk):
    """
    RF-18: mostra o que foi extraído. RF-21: permite corrigir tudo antes
    de confirmar. RF-14: permite "apenas analisar" sem vincular empresa.
    """
    documento = get_object_or_404(Documento, pk=pk)

    if request.method == "POST":
        if not pode_editar(request.user):
            messages.error(request, "Seu perfil (Consulta) não tem permissão para editar.")
            return redirect("revisar_documento", pk=pk)
        if "apenas_analisar" in request.POST:
            documento.revisado = True
            documento.save()
            messages.success(
                request, "Documento analisado sem vincular a uma empresa cadastrada."
            )
            return redirect("documento_list")

        form = DocumentoForm(request.POST, instance=documento)
        if form.is_valid():
            documento = form.save(commit=False)
            documento.revisado = True
            documento.save()
            messages.success(request, "Documento revisado e salvo com sucesso.")
            return redirect("documento_list")
    else:
        form = DocumentoForm(instance=documento)

    empresa_emitente = (
        Empresa.objects.filter(cnpj=documento.cnpj_emitente).first()
        if documento.cnpj_emitente else None
    )
    empresa_destinatario = (
        Empresa.objects.filter(cnpj=documento.cnpj_destinatario).first()
        if documento.cnpj_destinatario else None
    )

    fornecedor_emitente = (
    Fornecedor.objects.filter(cnpj=documento.cnpj_emitente).first()
        if documento.cnpj_emitente else None
        )
    return render(request, "core/revisar_documento.html", {
        "documento": documento,
        "form": form,
        "impostos": documento.dados_extraidos.get("impostos", []),
        "razao_social_destinatario": documento.dados_extraidos.get("razao_social_destinatario", ""),
        "empresa_emitente": empresa_emitente,
        "empresa_destinatario": empresa_destinatario,
        "fornecedor_emitente": fornecedor_emitente,
    })       
    
# ---------------------------------------------------------------------
# Gestão de usuários e papéis (RF-41, RF-42, RF-43)
# ---------------------------------------------------------------------
@requer_administrador
def usuario_list(request):
    usuarios = User.objects.select_related("perfil").order_by("username")
    return render(request, "core/usuario_list.html", {"usuarios": usuarios})


@requer_administrador
def novo_usuario(request):
    if request.method == "POST":
        form = NovoUsuarioForm(request.POST)
        if form.is_valid():
            usuario = form.save()
            usuario.perfil.papel = form.cleaned_data["papel"]
            usuario.perfil.save()
            messages.success(request, "Usuário criado com sucesso.")
            return redirect("usuario_list")
    else:
        form = NovoUsuarioForm()

    return render(request, "core/usuario_form.html", {"form": form, "titulo": "Novo Usuário"})


@requer_administrador
def editar_usuario(request, pk):
    usuario = get_object_or_404(User, pk=pk)

    if request.method == "POST":
        form = EditarUsuarioForm(request.POST, instance=usuario)
        if form.is_valid():
            form.save()
            messages.success(request, "Usuário atualizado com sucesso.")
            return redirect("usuario_list")
    else:
        form = EditarUsuarioForm(instance=usuario)

    return render(request, "core/usuario_form.html", {
        "form": form, "titulo": f"Editar — {usuario.username}"
    })    
    
# ---------------------------------------------------------------------
# Lançamentos contábeis (RF-33 a RF-36)
# ---------------------------------------------------------------------
@requer_edicao
def gerar_lancamento(request, documento_pk):
    """
    RF-33: gera um lançamento (em RASCUNHO) a partir dos dados já
    extraídos/revisados do documento. Só faz sentido se o documento já
    estiver vinculado a uma empresa administrada.
    """
    documento = get_object_or_404(Documento, pk=documento_pk)

    if not documento.empresa:
        messages.error(
            request,
            "Vincule esse documento a uma empresa cadastrada antes de gerar o lançamento."
        )
        return redirect("revisar_documento", pk=documento.pk)

    historico = f"Documento nº {documento.numero_documento}" if documento.numero_documento else "Lançamento gerado a partir de documento"
    if documento.fornecedor:
        historico += f" — Fornecedor: {documento.fornecedor.razao_social}"

    from django.utils import timezone

    lancamento = LancamentoContabil.objects.create(
        documento=documento,
        empresa=documento.empresa,
        data_lancamento=documento.data_emissao or timezone.localdate(),
        historico=historico,
        valor=documento.valor_total or 0,
    )

    messages.success(request, "Lançamento gerado como rascunho. Revise antes de confirmar.")
    return redirect("editar_lancamento", pk=lancamento.pk)


@login_required
def lancamento_list(request):
    """RF-36: histórico de todos os lançamentos, rascunhos e confirmados."""
    lancamentos = LancamentoContabil.objects.select_related(
        "empresa", "documento"
    ).order_by("-data_criacao")
    return render(request, "core/lancamento_list.html", {"lancamentos": lancamentos})


@requer_edicao
def editar_lancamento(request, pk):
    """RF-34: revisão/correção do lançamento enquanto ele é RASCUNHO."""
    lancamento = get_object_or_404(LancamentoContabil, pk=pk)

    if lancamento.status == LancamentoContabil.STATUS_CONFIRMADO:
        messages.info(request, "Esse lançamento já foi confirmado e não pode mais ser editado.")
        return redirect("lancamento_detail", pk=lancamento.pk)

    if request.method == "POST":
        form = LancamentoContabilForm(request.POST, instance=lancamento)
        if form.is_valid():
            form.save()
            messages.success(request, "Lançamento atualizado.")
            return redirect("lancamento_detail", pk=lancamento.pk)
    else:
        form = LancamentoContabilForm(instance=lancamento)

    return render(request, "core/lancamento_form.html", {
        "form": form, "lancamento": lancamento
    })


@login_required
def lancamento_detail(request, pk):
    """
    Mostra o lançamento. Se ainda for rascunho, oferece os botões de
    editar e confirmar (RF-35) — se já confirmado, é só leitura (RF-36).
    """
    lancamento = get_object_or_404(LancamentoContabil, pk=pk)
    return render(request, "core/lancamento_detail.html", {"lancamento": lancamento})


@requer_edicao
def confirmar_lancamento(request, pk):
    """RF-35: confirma o lançamento depois da revisão. Ação irreversível."""
    lancamento = get_object_or_404(LancamentoContabil, pk=pk)

    if lancamento.status == LancamentoContabil.STATUS_RASCUNHO:
        from django.utils import timezone
        lancamento.status = LancamentoContabil.STATUS_CONFIRMADO
        lancamento.data_confirmacao = timezone.now()
        lancamento.save()
        messages.success(request, "Lançamento confirmado com sucesso.")

    return redirect("lancamento_detail", pk=lancamento.pk)


class LancamentoDeleteView(RequerExclusaoMixin, DeleteView):
    """
    Só permite excluir rascunhos — um lançamento confirmado nunca pode
    ser apagado (RF-36: manter o histórico).
    """
    model = LancamentoContabil
    template_name = "core/confirm_delete.html"
    success_url = "/painel/lancamentos/"

    def get_queryset(self):
        return LancamentoContabil.objects.filter(status=LancamentoContabil.STATUS_RASCUNHO)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["cancel_url"] = self.success_url
        context["titulo"] = "Excluir lançamento (rascunho)"
        return context    
    
# ---------------------------------------------------------------------
# Certificados e configuração do sistema
# ---------------------------------------------------------------------
@login_required
def certificado_list(request):
    from .list_utils import aplicar_busca_e_ordenacao

    certificados, termo, ordenar = aplicar_busca_e_ordenacao(
        request,
        Certificado.objects.select_related("empresa"),
        campos_busca=["nome_empresa", "nome"],
        opcoes_ordenacao={"nome": "nome_empresa", "vencimento": "data_vencimento"},
        padrao="vencimento",
    )

    return render(request, "core/certificado_list.html", {
        "certificados": certificados, "termo": termo, "ordenar": ordenar,
    })

@requer_edicao
def novo_certificado(request):
    if request.method == "POST":
        form = CertificadoForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Certificado cadastrado com sucesso.")
            return redirect("certificado_list")
    else:
        form = CertificadoForm()

    return render(request, "core/certificado_form.html", {"form": form, "titulo": "Novo Certificado"})


@requer_edicao
def editar_certificado(request, pk):
    certificado = get_object_or_404(Certificado, pk=pk)

    if request.method == "POST":
        form = CertificadoForm(request.POST, instance=certificado)
        if form.is_valid():
            form.save()
            messages.success(request, "Certificado atualizado com sucesso.")
            return redirect("certificado_list")
    else:
        form = CertificadoForm(instance=certificado)

    return render(request, "core/certificado_form.html", {
        "form": form, "titulo": f"Editar — {certificado.nome_empresa}"
    })
    
@requer_edicao
def certificados_lote(request):
    """
    Carga rápida: várias linhas de nome + data de uma vez, igual à
    planilha que você já usa hoje.
    """
    if request.method == "POST":
        formset = CertificadoRapidoFormSet(request.POST, queryset=Certificado.objects.none())

        if formset.is_valid():
            criados = 0
            for form in formset:
                nome_empresa = form.cleaned_data.get("nome_empresa")
                if not nome_empresa:
                    continue
                certificado = form.save(commit=False)
                certificado.save()
                criados += 1

            messages.success(request, f"{criados} certificado(s) cadastrado(s) com sucesso.")
            return redirect("certificado_list")
    else:
        formset = CertificadoRapidoFormSet(queryset=Certificado.objects.none())

    return render(request, "core/certificados_lote.html", {"formset": formset})


@requer_edicao
def importar_certificados_excel(request):
    """
    Importa direto o Excel que você já mantém hoje. Linhas cuja data não
    dá pra interpretar com certeza entram sem data, para completar depois.
    """
    if request.method == "POST" and request.FILES.get("arquivo"):
        from .excel_utils import ler_planilha_certificados

        try:
            linhas = ler_planilha_certificados(request.FILES["arquivo"])
        except Exception:
            messages.error(request, "Não consegui ler esse arquivo. Confirme que é um .xlsx válido.")
            return redirect("importar_certificados_excel")

        criados = 0
        sem_data = 0

        for linha in linhas:
            Certificado.objects.create(
                nome_empresa=linha["nome_empresa"],
                data_vencimento=linha["data_vencimento"],
            )
            criados += 1
            if linha["data_vencimento"] is None:
                sem_data += 1

        if sem_data:
            messages.warning(
                request,
                f"{criados} certificado(s) importado(s), sendo {sem_data} sem "
                "data reconhecida — complete a data deles na lista."
            )
        else:
            messages.success(request, f"{criados} certificado(s) importado(s) com sucesso.")

        return redirect("certificado_list")

    return render(request, "core/importar_certificados_excel.html")

@requer_edicao
def importar_certificados_arquivo(request):
    """
    Importa certificado(s) direto do(s) arquivo(s) .pfx/.p12. Se já
    existir um certificado do mesmo titular com a MESMA data de
    vencimento, é um duplicado exato. Se existir um do mesmo titular
    com data DIFERENTE, é uma atualização (renovação) — nos dois casos
    fica pendente para você decidir na tela seguinte.
    """
    if request.method == "POST":
        from django.core.files.base import ContentFile
        from django.core.files.storage import default_storage
        from .cert_utils import extrair_dados_certificado, eh_arquivo_certificado
        from .cert_dedupe import buscar_certificado_relacionado, extrair_cnpj_do_texto

        senha = request.POST.get("senha", "")
        arquivos = request.FILES.getlist("arquivos")
        candidatos = [a for a in arquivos if eh_arquivo_certificado(a.name)]
        ignorados = len(arquivos) - len(candidatos)

        importados = 0
        falhas_senha = []
        pendentes = request.session.get("certificados_pendentes_revisao", [])

        for arquivo in candidatos:
            conteudo = arquivo.read()

            try:
                dados = extrair_dados_certificado(conteudo, senha)
            except Exception:
                falhas_senha.append(arquivo.name)
                continue

            cnpj = extrair_cnpj_do_texto(dados["nome_completo_cn"])
            tipo, existente = buscar_certificado_relacionado(
                dados["nome_empresa"], dados["data_vencimento"], cnpj
            )

            if tipo:
                caminho_temp = default_storage.save(
                    f"certificados/_pendentes/{arquivo.name}", ContentFile(conteudo)
                )
                pendentes.append({
                    "tipo": tipo,
                    "caminho": caminho_temp,
                    "nome_arquivo": arquivo.name,
                    "nome_empresa": dados["nome_empresa"],
                    "data_vencimento": str(dados["data_vencimento"]),
                    "nome_completo_cn": dados["nome_completo_cn"],
                    "existente_id": existente.id,
                })
                continue

            certificado = Certificado(
                nome_empresa=dados["nome_empresa"],
                data_vencimento=dados["data_vencimento"],
                observacao=dados["nome_completo_cn"],
            )
            certificado.arquivo.save(arquivo.name, ContentFile(conteudo), save=False)
            certificado.save()
            importados += 1

        request.session["certificados_pendentes_revisao"] = pendentes
        total_agora = Certificado.objects.count()

        if importados:
            messages.success(request, f"{importados} certificado(s) novo(s) importado(s). Total agora: {total_agora}.")

        if ignorados:
            messages.info(request, f"{ignorados} arquivo(s) ignorados por não serem .pfx/.p12.")

        if falhas_senha:
            messages.warning(request, "Não foi possível ler (senha incorreta): " + ", ".join(falhas_senha))

        if pendentes:
            messages.warning(
                request,
                f"{len(pendentes)} arquivo(s) precisam da sua decisão (duplicado ou atualização)."
            )
            return redirect("resolver_duplicados_certificados")

        if not importados and not ignorados and not falhas_senha:
            messages.info(request, f"Nada novo para importar. Total agora: {total_agora}.")

        return redirect("certificado_list")

    return render(request, "core/importar_certificados_arquivo.html")


@requer_edicao
def resolver_duplicados_certificados(request):
    """
    Para "exato": Substituir (troca o arquivo/dados do existente) ou
    Cancelar (descarta o novo). Para "atualizacao": Atualizar (o novo
    vencimento substitui o antigo no MESMO registro — não deixa dois),
    Manter os dois (cria um registro separado mesmo assim), ou Cancelar.
    """
    from django.core.files.base import ContentFile
    from django.core.files.storage import default_storage

    pendentes = request.session.get("certificados_pendentes_revisao", [])

    if not pendentes:
        messages.info(request, "Não há certificados pendentes de decisão.")
        return redirect("certificado_list")

    if request.method == "POST":
        atualizados = 0
        substituidos = 0
        mantidos_separados = 0
        cancelados = 0
        restantes = []

        for i, item in enumerate(pendentes):
            acao = request.POST.get(f"acao_{i}")

            if acao in ("substituir", "atualizar"):
                with default_storage.open(item["caminho"], "rb") as f:
                    conteudo = f.read()

                existente = get_object_or_404(Certificado, pk=item["existente_id"])
                existente.nome_empresa = item["nome_empresa"]
                existente.data_vencimento = item["data_vencimento"]
                existente.observacao = item["nome_completo_cn"]
                existente.arquivo.save(item["nome_arquivo"], ContentFile(conteudo), save=False)
                existente.save()

                default_storage.delete(item["caminho"])
                if acao == "substituir":
                    substituidos += 1
                else:
                    atualizados += 1

            elif acao == "manter_separado":
                with default_storage.open(item["caminho"], "rb") as f:
                    conteudo = f.read()

                novo = Certificado(
                    nome_empresa=item["nome_empresa"],
                    data_vencimento=item["data_vencimento"],
                    observacao=item["nome_completo_cn"],
                )
                novo.arquivo.save(item["nome_arquivo"], ContentFile(conteudo), save=False)
                novo.save()

                default_storage.delete(item["caminho"])
                mantidos_separados += 1

            elif acao == "cancelar":
                default_storage.delete(item["caminho"])
                cancelados += 1

            else:
                restantes.append(item)

        request.session["certificados_pendentes_revisao"] = restantes

        total_agora = Certificado.objects.count()
        messages.success(
            request,
            f"{atualizados} atualizado(s), {substituidos} substituído(s), "
            f"{mantidos_separados} mantido(s) separado(s), {cancelados} descartado(s). "
            f"Total agora: {total_agora}."
        )

        if not restantes:
            return redirect("certificado_list")
        return redirect("resolver_duplicados_certificados")

    contexto = []
    for i, item in enumerate(pendentes):
        existente = Certificado.objects.filter(pk=item["existente_id"]).first()
        contexto.append({"indice": i, "novo": item, "existente": existente, "tipo": item["tipo"]})

    return render(request, "core/resolver_duplicados_certificados.html", {"itens": contexto})


@requer_exclusao
def excluir_certificados_selecionados(request):
    """Exclui em massa os certificados marcados na lista (checkboxes)."""
    if request.method == "POST":
        ids = request.POST.getlist("selecionados")
        if ids:
            Certificado.objects.filter(id__in=ids).delete()
            messages.success(
                request,
                f"{len(ids)} certificado(s) excluído(s). Total agora: {Certificado.objects.count()}."
            )
        else:
            messages.info(request, "Nenhum certificado selecionado.")

    return redirect("certificado_list")

@requer_edicao
def retentar_certificados(request):
    """
    Segunda chance para os certificados que não abriram com a senha
    informada na tela anterior — cada um pode receber sua própria
    senha aqui, ou ser descartado, sem precisar selecionar o arquivo
    de novo (ele já está guardado temporariamente).
    """
    from django.core.files.base import ContentFile
    from django.core.files.storage import default_storage
    from .cert_utils import extrair_dados_certificado

    pendentes = request.session.get("certificados_pendentes", [])

    if not pendentes:
        messages.info(request, "Não há certificados pendentes de senha no momento.")
        return redirect("certificado_list")

    if request.method == "POST":
        acao = request.POST.get("acao")

        if acao == "descartar":
            indice = int(request.POST.get("indice"))
            if 0 <= indice < len(pendentes):
                item = pendentes.pop(indice)
                default_storage.delete(item["caminho"])
                messages.info(request, f'Arquivo "{item["nome_arquivo"]}" descartado.')

        else:
            novos_pendentes = []
            importados = 0

            for i, item in enumerate(pendentes):
                senha = request.POST.get(f"senha_{i}", "")

                with default_storage.open(item["caminho"], "rb") as f:
                    conteudo = f.read()

                try:
                    dados = extrair_dados_certificado(conteudo, senha)
                except Exception:
                    novos_pendentes.append(item)
                    continue

                certificado = Certificado(
                    nome_empresa=dados["nome_empresa"],
                    data_vencimento=dados["data_vencimento"],
                    observacao=dados["nome_completo_cn"],
                )
                certificado.arquivo.save(item["nome_arquivo"], ContentFile(conteudo), save=False)
                certificado.save()
                default_storage.delete(item["caminho"])
                importados += 1

            pendentes = novos_pendentes

            if importados:
                messages.success(request, f"{importados} certificado(s) importado(s) com sucesso.")
            if pendentes:
                messages.warning(request, f"Ainda restam {len(pendentes)} certificado(s) com senha incorreta.")

        request.session["certificados_pendentes"] = pendentes

        if not pendentes:
            return redirect("certificado_list")
        return redirect("retentar_certificados")

    return render(request, "core/retentar_certificados.html", {
        "pendentes": list(enumerate(pendentes))
    })
    
class CertificadoDeleteView(RequerExclusaoMixin, DeleteView):
    model = Certificado
    template_name = "core/confirm_delete.html"
    success_url = "/painel/certificados/"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["cancel_url"] = self.success_url
        context["titulo"] = "Excluir certificado"
        return context


@login_required
def exportar_certificados_excel(request):
    """
    Exporta os certificados para .xlsx usando o mesmo semáforo (iconSet
    de 3 cores) da tela — vencido/próximo/em dia.
    """
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.formatting.rule import IconSetRule
    from django.http import HttpResponse

    workbook = openpyxl.Workbook()
    aba = workbook.active
    aba.title = "Certificados"

    cabecalho_fill = PatternFill("solid", fgColor="E7E7E7")
    centro = Alignment(horizontal="center")

    colunas = ["Empresa", "Certificado", "Vencimento", "Status (semáforo)"]
    for col, titulo in enumerate(colunas, start=1):
        celula = aba.cell(row=1, column=col, value=titulo)
        celula.font = Font(bold=True)
        celula.fill = cabecalho_fill

    linha = 2
    for certificado in Certificado.objects.select_related("empresa").order_by("data_vencimento"):
        aba.cell(row=linha, column=1, value=certificado.nome_empresa)
        aba.cell(row=linha, column=2, value=certificado.nome)
        aba.cell(
            row=linha, column=3,
            value=certificado.data_vencimento.strftime("%d/%m/%Y") if certificado.data_vencimento else ""
        )

        celula_status = aba.cell(row=linha, column=4)
        if certificado.data_vencimento:
            celula_status.value = certificado.status_calculado()
        celula_status.alignment = centro

        linha += 1

    if linha > 2:
        aba.conditional_formatting.add(
            f"D2:D{linha - 1}",
            IconSetRule(icon_style="3TrafficLights1", type="num", values=[0, 2, 3], showValue=False)
        )

    aba.column_dimensions["A"].width = 35
    aba.column_dimensions["B"].width = 35
    aba.column_dimensions["C"].width = 15
    aba.column_dimensions["D"].width = 18

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="certificados.xlsx"'
    workbook.save(response)
    return response


@requer_administrador
def configuracao_sistema(request):
    config = ConfiguracaoSistema.obter()

    if request.method == "POST":
        form = ConfiguracaoSistemaForm(request.POST, instance=config)
        if form.is_valid():
            form.save()
            messages.success(request, "Configurações atualizadas com sucesso.")
            return redirect("configuracao_sistema")
    else:
        form = ConfiguracaoSistemaForm(instance=config)

    return render(request, "core/configuracao_sistema.html", {"form": form})    

@requer_edicao
def importar_rotinas_excel(request):
    from datetime import date

    if request.method == "POST":
        form = ImportarRotinasForm(request.POST, request.FILES)
        if form.is_valid():
            from .rotinas_import_utils import processar_planilha, importar_log_para_banco

            try:
                log = processar_planilha(request.FILES["arquivo"], ano_padrao=form.cleaned_data["ano"])
            except Exception:
                messages.error(
                    request,
                    "Não consegui ler esse arquivo. Confirme que é um .xlsx no mesmo formato do modelo."
                )
                return redirect("importar_rotinas_excel")

            resultado = importar_log_para_banco(log)

            messages.success(
                request,
                f"Importação concluída: {resultado['empresas_criadas']} empresa(s) nova(s), "
                f"{resultado['empresas_existentes']} já existente(s), "
                f"{resultado['obrigacoes_criadas']} obrigação(ões) nova(s), "
                f"{resultado['cumprimentos_gravados']} status gravados."
            )
            return redirect("empresa_list")
    else:
        form = ImportarRotinasForm(initial={"ano": date.today().year})

    return render(request, "core/importar_rotinas_excel.html", {"form": form})

# ---------------------------------------------------------------------
# --------------------------------------------------------------------
#      VENCIMIENTOS
# ---------------------------------------------------------------------
# ---------------------------------------------------------------------

@login_required
def vencimentos_list(request):
    """
    Painel de vencimentos agrupado por empresa (obrigações e impostos
    juntos, debaixo de cada uma), separado por regime tributário em
    abas. Filtros: status (todas/atrasada/risco/em dia), busca por
    texto, e ordenação (mais próximo do vencimento ou alfabética).
    """
    termo = request.GET.get("q", "").strip().lower()
    ordenar = request.GET.get("ordenar", "vencimento")
    filtro_status = request.GET.get("status", "todas")
    aba_selecionada = request.GET.get("aba", "")

    def status_bate(status_num):
        return filtro_status == "todas" or str(status_num) == filtro_status

    cumprimentos = CumprimentoObrigacao.objects.filter(cumprido=False).select_related(
        "competencia__empresa__regime_tributario", "obrigacao"
    )
    impostos_qs = ApuracaoImposto.objects.filter(
        status__in=["PENDENTE", "ATRASADO"]
    ).select_related("competencia__empresa__regime_tributario", "imposto")

    por_regime = {regime.nome: {} for regime in RegimeTributario.objects.all()}

    for cumprimento in cumprimentos:
        status = cumprimento.status_calculado()
        if not status_bate(status):
            continue

        empresa = cumprimento.competencia.empresa
        if termo and termo not in empresa.razao_social.lower() and termo not in cumprimento.obrigacao.nome.lower():
            continue

        regime_nome = empresa.regime_tributario.nome if empresa.regime_tributario else "Sem regime definido"
        empresas_dict = por_regime.setdefault(regime_nome, {})
        empresas_dict.setdefault(empresa.razao_social, []).append({
            "tipo": "obrigacao",
            "nome": cumprimento.obrigacao.nome,
            "vencimento": cumprimento.data_vencimento_calculada(),
            "status": status,
            "pk": cumprimento.pk,
        })

    for apuracao in impostos_qs:
        status = 1 if apuracao.status == "ATRASADO" else 3
        if not status_bate(status):
            continue

        empresa = apuracao.competencia.empresa
        if termo and termo not in empresa.razao_social.lower() and termo not in apuracao.imposto.nome.lower():
            continue

        regime_nome = empresa.regime_tributario.nome if empresa.regime_tributario else "Sem regime definido"
        empresas_dict = por_regime.setdefault(regime_nome, {})
        empresas_dict.setdefault(empresa.razao_social, []).append({
            "tipo": "imposto",
            "nome": apuracao.imposto.nome,
            "vencimento": apuracao.data_vencimento,
            "status": status,
            "pk": apuracao.pk,
        })

    abas = []
    for i, (regime_nome, empresas_dict) in enumerate(por_regime.items()):
        lista_empresas = []
        for nome_empresa, itens in empresas_dict.items():
            itens_validos = [it for it in itens if it["vencimento"]]
            if not itens_validos:
                continue
            itens_validos.sort(key=lambda it: it["vencimento"])
            lista_empresas.append({
                "nome": nome_empresa,
                "itens": itens_validos,
                "mais_proximo": itens_validos[0]["vencimento"],
            })

        if ordenar == "nome":
            lista_empresas.sort(key=lambda e: e["nome"].lower())
        else:
            lista_empresas.sort(key=lambda e: e["mais_proximo"])

        abas.append({"chave": f"regime-{i}", "titulo": regime_nome, "empresas": lista_empresas})

    chaves_abas = {a["chave"] for a in abas}
    if aba_selecionada not in chaves_abas:
        aba_selecionada = abas[0]["chave"] if abas else ""
    for aba in abas:
        aba["ativa"] = aba["chave"] == aba_selecionada

    return render(request, "core/vencimentos_list.html", {
        "abas": abas,
        "termo": termo,
        "ordenar": ordenar,
        "filtro_status": filtro_status,
        "aba_selecionada": aba_selecionada,
    })

@requer_edicao
def atualizar_vencimento_cumprimento(request, pk):
    """
    Guarda a DATA COMPLETA de vencimento escolhida por você para essa
    obrigação, nesse mês específico — não um dia recorrente automático.
    """
    from datetime import datetime

    cumprimento = get_object_or_404(CumprimentoObrigacao, pk=pk)

    if request.method == "POST":
        valor = request.POST.get("data_vencimento")
        try:
            nova_data = datetime.strptime(valor, "%Y-%m-%d").date()
            cumprimento.data_vencimento_manual = nova_data
            cumprimento.save()
            messages.success(request, "Data de vencimento atualizada.")
        except (TypeError, ValueError):
            messages.error(request, "Data inválida.")

    aba = request.POST.get("aba", "")
    if aba:
        return redirect(f"/painel/vencimentos/?aba={aba}")
    return redirect(request.META.get("HTTP_REFERER") or "vencimentos_list")

@login_required
def notificacoes_json(request):
    """
    La campanita llama a esta URL cada cierto tiempo (JS). Genera
    notificações nuevas como mucho cada 15 minutos (no en cada llamada,
    para no recorrer toda la base de datos todo el rato).
    """
    from datetime import timedelta, datetime
    from django.utils import timezone
    from django.http import JsonResponse
    from .notificacoes import gerar_notificacoes

    agora = timezone.now()
    ultima = request.session.get("ultima_geracao_notificacoes")

    deve_gerar = True
    if ultima:
        deve_gerar = (agora - datetime.fromisoformat(ultima)) > timedelta(minutes=15)

    if deve_gerar:
        gerar_notificacoes()
        request.session["ultima_geracao_notificacoes"] = agora.isoformat()

    nao_lidas = Notificacao.objects.filter(lida=False)[:20]

    return JsonResponse({
        "total_nao_lidas": Notificacao.objects.filter(lida=False).count(),
        "notificacoes": [
            {"id": n.id, "mensagem": n.mensagem, "link": n.link, "tipo": n.tipo}
            for n in nao_lidas
        ],
    })


@login_required
def notificacoes_list(request):
    notificacoes = Notificacao.objects.all()[:200]
    return render(request, "core/notificacoes_list.html", {"notificacoes": notificacoes})


@login_required
def marcar_notificacao_lida(request, pk):
    notificacao = get_object_or_404(Notificacao, pk=pk)
    notificacao.lida = True
    notificacao.save()
    return redirect(notificacao.link or "notificacoes_list")


@login_required
def marcar_todas_lidas(request):
    Notificacao.objects.filter(lida=False).update(lida=True)
    messages.success(request, "Todas as notificações foram marcadas como lidas.")
    return redirect("notificacoes_list")
 